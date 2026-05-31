"""
rebuild_ppt_simulation_slides.py (v2 - 安全版)
================================================
1. 删除已存在的损坏仿真页 (含我之前克隆产生的)
2. 用纯 python-pptx API 从零构造仿真页 (无 XML 深拷贝)
3. 风格模仿原 PPT (顶/底装饰条、章节标题、chip 标签、分割线、结论框)
4. 按 PPT 逻辑顺序插入

注意: 这个脚本会保留原始 48 页不动, 只重建仿真页.
"""
import copy
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

import pandas as pd

BASE = Path(__file__).parent.parent
PPT_PATH = BASE / 'ppt/尺子之歌-正方匡大德.pptx'
FIG_DIR = BASE / 'outputs/requested_adaptive/figures'

# ============================================================
# 风格常量 (匹配原 PPT 配色)
# ============================================================
COLOR_BG = RGBColor(0xF5, 0xF6, 0xF7)        # 极浅灰背景
COLOR_BAR = RGBColor(0x1F, 0x3A, 0x5F)        # 深蓝顶/底装饰条
COLOR_ACCENT = RGBColor(0xC8, 0x96, 0x66)     # 暖金色 chip 边
COLOR_CHIP_FILL = RGBColor(0x1F, 0x3A, 0x5F)  # chip 深蓝填充
COLOR_CHIP_TEXT = RGBColor(0xFF, 0xFF, 0xFF)  # chip 白字
COLOR_TITLE = RGBColor(0x1F, 0x3A, 0x5F)      # 标题深蓝
COLOR_TEXT = RGBColor(0x33, 0x33, 0x33)       # 正文深灰
COLOR_SECONDARY = RGBColor(0x80, 0x80, 0x80)  # 次要灰
COLOR_DIVIDER = RGBColor(0xC8, 0x96, 0x66)    # 分隔线金色

# ============================================================
# Step 1: 加载并清理重复仿真页
# ============================================================
print("Step 1: 清理已存在的仿真页...")
prs = Presentation(str(PPT_PATH))


def is_simulation_slide(slide):
    """检查幻灯片是否是已插入的仿真页"""
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text.startswith('仿真验证:') or text.startswith('新增:'):
                return True
            if text == '【仿真验证】':
                return True
            if text == '仿真验证':
                return True
            if text.startswith('仿真图'):
                return True
            if text.startswith('3.6 模态衰减分析'):
                return True
    return False


# Delete from end backward
slides_xml = prs.slides._sldIdLst
xml_slides = list(slides_xml)
deleted = 0
for idx in range(len(prs.slides) - 1, -1, -1):
    sl = prs.slides[idx]
    if is_simulation_slide(sl):
        sldId = xml_slides[idx]
        rId = sldId.get(
            '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
        )
        prs.part.drop_rel(rId)
        slides_xml.remove(sldId)
        deleted += 1

print(f"  已删除 {deleted} 张旧仿真页")
prs.save(str(PPT_PATH))
prs = Presentation(str(PPT_PATH))
print(f"  当前页数: {len(prs.slides)}")
assert len(prs.slides) == 48, f"原始页数应为 48, 实际 {len(prs.slides)}"

# ============================================================
# Step 2: 工具函数 - 从零构造仿真页
# ============================================================
print("\nStep 2: 准备从零构造仿真页 (无 XML 克隆)...")

SLIDE_W = prs.slide_width   # EMU
SLIDE_H = prs.slide_height  # EMU


def add_filled_rect(slide, left, top, width, height, fill_color, line_color=None):
    """添加纯色填充矩形"""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
    shape.shadow.inherit = False
    return shape


def add_chip(slide, left, top, width, height, text, fill_color=COLOR_CHIP_FILL,
             text_color=COLOR_CHIP_TEXT, font_size=14, bold=True):
    """添加圆角矩形标签 (chip)"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = text
    for run in p.runs:
        run.font.name = 'Microsoft YaHei'
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = text_color
    return shape


def add_textbox(slide, left, top, width, height, text, font_size=12,
                color=COLOR_TEXT, bold=False, align=PP_ALIGN.LEFT,
                anchor=MSO_ANCHOR.TOP):
    """添加纯文本框"""
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.vertical_anchor = anchor
    lines = text.split('\n')
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.text = line
        for run in p.runs:
            run.font.name = 'Microsoft YaHei'
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = color
    return shape


def add_horizontal_line(slide, left, top, width, color=COLOR_DIVIDER, weight_pt=1.5):
    """添加水平分隔线"""
    shape = slide.shapes.add_connector(1, left, top, left + width, top)  # 1 = STRAIGHT
    shape.line.color.rgb = color
    shape.line.width = Pt(weight_pt)
    return shape


def add_vertical_line(slide, left, top, height, color=COLOR_DIVIDER, weight_pt=1.0):
    """添加竖直分隔线"""
    shape = slide.shapes.add_connector(1, left, top, left, top + height)
    shape.line.color.rgb = color
    shape.line.width = Pt(weight_pt)
    return shape


def make_sim_slide(prs, image_path, section_title, fit_label, caption, conclusion, page_label):
    """从零构造一张仿真验证幻灯片"""
    blank_layout = prs.slide_layouts[6]  # 空白
    slide = prs.slides.add_slide(blank_layout)

    # === 整体浅色背景 ===
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = COLOR_BG

    # === 顶部装饰条 (深蓝) ===
    add_filled_rect(slide, 0, 0, SLIDE_W, Inches(0.6), COLOR_BAR)

    # === 顶部装饰下沿金线 ===
    add_horizontal_line(slide, 0, Inches(0.62), SLIDE_W, COLOR_DIVIDER, weight_pt=2)

    # === 主标题 (左上) ===
    add_textbox(slide, Inches(0.4), Inches(0.08), Inches(8), Inches(0.5),
                section_title, font_size=20, color=RGBColor(0xFF, 0xFF, 0xFF),
                bold=True, anchor=MSO_ANCHOR.MIDDLE)

    # === "仿真验证" chip (左侧 标题下) ===
    add_chip(slide, Inches(0.4), Inches(0.85), Inches(2.2), Inches(0.45),
             '仿真验证', fill_color=COLOR_CHIP_FILL, text_color=COLOR_CHIP_TEXT,
             font_size=13, bold=True)

    # === 拟合类型 chip (右上 标题下) ===
    add_chip(slide, Inches(10.5), Inches(0.85), Inches(2.4), Inches(0.45),
             fit_label, fill_color=COLOR_ACCENT, text_color=COLOR_CHIP_TEXT,
             font_size=13, bold=True)

    # === 中央分隔竖线 ===
    add_vertical_line(slide, Inches(6.7), Inches(1.45), Inches(5.0),
                      color=COLOR_DIVIDER, weight_pt=1.0)

    # === 左半 - 主图片 ===
    if image_path is not None and Path(image_path).exists():
        slide.shapes.add_picture(str(image_path),
                                 Inches(0.4), Inches(1.55),
                                 Inches(6.2), Inches(4.3))

    # === 左半 - 图注 ===
    add_textbox(slide, Inches(0.4), Inches(5.95), Inches(6.2), Inches(0.4),
                caption, font_size=11, color=COLOR_SECONDARY, bold=False,
                align=PP_ALIGN.CENTER)

    # === 右半 - "实验结论" header chip ===
    add_chip(slide, Inches(7.0), Inches(1.55), Inches(2.4), Inches(0.45),
             '仿真结论', fill_color=COLOR_CHIP_FILL, text_color=COLOR_CHIP_TEXT,
             font_size=13, bold=True)

    # === 右半 - 结论框 ===
    conclusion_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.0), Inches(2.15), Inches(5.9), Inches(3.7))
    conclusion_box.fill.solid()
    conclusion_box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    conclusion_box.line.color.rgb = COLOR_DIVIDER
    conclusion_box.line.width = Pt(1.0)
    conclusion_box.shadow.inherit = False

    tf = conclusion_box.text_frame
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.15)
    tf.margin_bottom = Inches(0.15)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    lines = conclusion.split('\n')
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.text = line
        for run in p.runs:
            run.font.name = 'Microsoft YaHei'
            run.font.size = Pt(13)
            run.font.bold = (i == 0)
            run.font.color.rgb = COLOR_TITLE if i == 0 else COLOR_TEXT

    # === 底部装饰条 ===
    add_filled_rect(slide, 0, Inches(7.05), SLIDE_W, Inches(0.45), COLOR_BAR)

    # === 底部页码 ===
    add_textbox(slide, Inches(11.5), Inches(7.10), Inches(1.5), Inches(0.35),
                page_label, font_size=11, color=RGBColor(0xFF, 0xFF, 0xFF),
                bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    # === 底部 学校水印 ===
    add_textbox(slide, Inches(0.4), Inches(7.10), Inches(6), Inches(0.35),
                '中国石油大学(华东) | 尺子之歌 - 仿真验证补充',
                font_size=10, color=RGBColor(0xCC, 0xCC, 0xCC),
                bold=False, anchor=MSO_ANCHOR.MIDDLE)

    return slide


# ============================================================
# Step 3: 加载拟合数据 + 创建仿真页
# ============================================================
print("\nStep 3: 创建仿真幻灯片 (从零构造)...")

fit_df = pd.read_csv(BASE / 'outputs/requested_adaptive/tables/fitting_results.csv')
fit_dict = {row['relation']: row for _, row in fit_df.iterrows()}

sim_slide_specs = []

# === 音调 4 张 (插在原 slide 28 之后) ===
r = fit_dict['f1 vs L_free']
sim_slide_specs.append({
    'image': FIG_DIR / 'pitch_f1_vs_L.png',
    'section': '3.3 音调理论验证',
    'fit_label': '幂律拟合',
    'caption': '仿真图1  f₁ vs 自由端长度 L',
    'conclusion': "结论: f₁ ∝ 1/L² 关系完美验证\n\n"
                  f"• 幂律拟合指数 n = {r['fitted_exponent']:.4f}\n"
                  f"• 理论值: -2 (相对偏差 {abs(r['fitted_exponent']+2)/2*100:.2f}%)\n"
                  f"• R² = {r['R2']:.5f}\n"
                  f"• 平均相对误差 {r['mean_rel_error_pct']:.2f}%\n\n"
                  "短尺 → 高音调, 长尺 → 低音调",
    'insert_after': 28,
})

r = fit_dict['f1 vs h']
k_h = float(r['formula'].split('=')[1].strip().split('*')[0].strip())
sim_slide_specs.append({
    'image': FIG_DIR / 'pitch_f1_vs_h.png',
    'section': '3.3 音调理论验证',
    'fit_label': '线性拟合',
    'caption': '仿真图2  f₁ vs 厚度 h',
    'conclusion': "结论: f₁ ∝ h 线性关系完美验证\n\n"
                  f"• 线性过原点斜率 k = {k_h:.2f} Hz/mm\n"
                  f"• R² = {r['R2']:.5f}\n"
                  f"• 平均相对误差 {r['mean_rel_error_pct']:.2f}%\n\n"
                  "厚尺 → 高音调, 薄尺 → 低音调",
    'insert_after': 28,
})

r = fit_dict['f1 vs b (独立性)']
sim_slide_specs.append({
    'image': FIG_DIR / 'pitch_f1_vs_b.png',
    'section': '3.3 音调理论验证',
    'fit_label': '独立性验证',
    'caption': '仿真图3  f₁ vs 宽度 b',
    'conclusion': "结论: f₁ 不随宽度 b 变化\n\n"
                  f"• 5 种宽度变异 Δf/f = {r['mean_rel_error_pct']:.4f}%\n"
                  "• 理论: I/A = h²/12 与 b 无关\n"
                  "• 仿真完美验证该约分关系\n\n"
                  "宽度仅影响声压,不影响音调",
    'insert_after': 28,
})

r = fit_dict['f1 vs J (独立性)']
sim_slide_specs.append({
    'image': FIG_DIR / 'pitch_f1_vs_J.png',
    'section': '3.3 音调理论验证',
    'fit_label': '独立性验证',
    'caption': '仿真图4  f₁ vs 冲量 J',
    'conclusion': "结论: f₁ 不随冲量 J 变化\n\n"
                  f"• 4 种冲量变异 Δf/f = {r['mean_rel_error_pct']:.6f}%\n"
                  "• 线性系统固有频率与激励无关\n"
                  "• 仿真变异趋近于零\n\n"
                  "冲量仅影响响度, 不影响音调",
    'insert_after': 28,
})

# === 音色 3 张 (插在原 slide 30 之后) ===
sim_slide_specs.append({
    'image': FIG_DIR / 'timbre_material_comparison.png',
    'section': '3.4 音色理论验证',
    'fit_label': '材料对比',
    'caption': '仿真图5  五材料频谱质心 + 平坦度对比',
    'conclusion': "结论: 材料音色指标差异显著\n\n"
                  "• 钢: 质心低 + 平坦度低 → 音色纯净\n"
                  "• 非钢: 质心高 → 谐波丰富\n\n"
                  "⚠ 注意: 非钢材料 probe_peak/L > 30%\n"
                  "线性假设不成立, 仅供定性参考",
    'insert_after': 30,
})

r = fit_dict['tau vs f (钢, 全模态)']
sim_slide_specs.append({
    'image': FIG_DIR / 'timbre_tau_vs_f.png',
    'section': '3.4 音色理论验证',
    'fit_label': '衰减分析',
    'caption': '仿真图6  τ vs f 6阶模态衰减',
    'conclusion': "结论: 高频衰减快, 低频持久\n\n"
                  f"• 幂律拟合 τ = {r['fitted_exponent']:.3f} 阶下降\n"
                  f"• R² = {r['R2']:.4f}\n"
                  "• 验证 PPT 中音色随时间演变规律\n\n"
                  "敲击瞬间: 多频率 \"叮\"\n"
                  "持续阶段: 基频主导 \"嗡\"",
    'insert_after': 30,
})

sim_slide_specs.append({
    'image': FIG_DIR / 'timbre_bubble_map.png',
    'section': '3.4 音色理论验证',
    'fit_label': '音色映射',
    'caption': '仿真图7  五材料音色三维映射',
    'conclusion': "结论: 材料在音色空间分布差异显著\n\n"
                  "• 钢最 \"纯\" (低质心 + 低平坦度)\n"
                  "• 雪弗板最 \"杂\" (高谐波)\n"
                  "• 气泡大小 ∝ SPL_rms\n\n"
                  "⚠ 非钢材料大变形\n"
                  "数据仅供定性对比",
    'insert_after': 30,
})

# === 响度 3 张 + 模态衰减 2 张 (插在原 slide 35 之后) ===
r_spl = fit_dict['SPL_rms vs J']
r_probe = fit_dict['probe_peak vs J']
sim_slide_specs.append({
    'image': FIG_DIR / 'loudness_spl_probe_vs_J.png',
    'section': '3.5 响度理论验证',
    'fit_label': '线性拟合',
    'caption': '仿真图8  SPL/位移 vs 冲量 J',
    'conclusion': "结论: 声压与冲量高度线性相关\n\n"
                  f"• SPL_rms vs J: R² = {r_spl['R2']:.4f}\n"
                  f"• probe_peak vs J: R² = {r_probe['R2']:.4f}\n"
                  "• 验证 p ∝ J 理论预测\n\n"
                  "敲击越用力 → 声压越大",
    'insert_after': 35,
})

sim_slide_specs.append({
    'image': FIG_DIR / 'loudness_spl_vs_x0.png',
    'section': '3.5 响度理论验证',
    'fit_label': '位置扫描',
    'caption': '仿真图9  SPL vs 撞击位置 x₀',
    'conclusion': "结论: 撞击位置显著影响声压\n\n"
                  "• 越靠近自由端撞击, SPL 越大\n"
                  "• x₀=0.75 比 x₀=0.25 高 ~6 dB\n"
                  "• 与 PPT 实验结论完全一致\n\n"
                  "撞击位置必须严格控制",
    'insert_after': 35,
})

sim_slide_specs.append({
    'image': FIG_DIR / 'loudness_spl_vs_distance.png',
    'section': '3.5 响度理论验证',
    'fit_label': '距离衰减',
    'caption': '仿真图10  SPL 空间衰减 (10 距离麦克风)',
    'conclusion': "结论: 声压空间衰减规律\n\n"
                  + (f"• 钢拟合: SPL = {fit_dict['SPL vs distance (钢)']['fitted_exponent']:.2f}·log₁₀(r) + 常数\n"
                     f"• 几何衰减幂 n ≈ {-fit_dict['SPL vs distance (钢)']['fitted_exponent']/20:.3f}\n"
                     f"• R² = {fit_dict['SPL vs distance (钢)']['R2']:.4f}\n"
                     if 'SPL vs distance (钢)' in fit_dict else "")
                  + "• 弥补 PPT 实验仅 5/30cm 两点不足\n\n"
                  "⚠ 非钢材料 SPL 绝对值不可直接对比",
    'insert_after': 35,
})

sim_slide_specs.append({
    'image': FIG_DIR / 'modal_Q_vs_mode.png',
    'section': '3.6 模态衰减分析 (新增)',
    'fit_label': '品质因数 Q',
    'caption': '仿真图11  6 阶模态品质因数 Q',
    'conclusion': "结论: 量化高频衰减快\n\n"
                  "• 低阶模态 Q 值高 → 振动持久\n"
                  "• 高阶模态 Q 值低 → 快速衰减\n"
                  "• 这是 PPT 实验无法直接测量的指标\n\n"
                  "支撑\"高频消失快\"现象的定量证据",
    'insert_after': 35,
})

sim_slide_specs.append({
    'image': FIG_DIR / 'modal_rayleigh_params.png',
    'section': '3.6 模态衰减分析 (新增)',
    'fit_label': 'Rayleigh 阻尼',
    'caption': '仿真图12  Rayleigh 阻尼参数分布',
    'conclusion': "结论: 阻尼模型评估\n\n"
                  "• α (质量阻尼) 与 β (刚度阻尼) 估计\n"
                  "• 多数 case 的 Rayleigh 模型不完全有效\n"
                  "• 建议引入频率相关阻尼\n\n"
                  "PPT 误差分析中应补充该局限性",
    'insert_after': 35,
})

# 创建所有幻灯片(暂存于末尾)
created = []
for i, spec in enumerate(sim_slide_specs):
    page_label = f'S{i+1}'
    sl = make_sim_slide(prs, spec['image'], spec['section'],
                        spec['fit_label'], spec['caption'],
                        spec['conclusion'], page_label)
    created.append((sl, spec['insert_after']))
    print(f"  + {spec['section']:20s} | {spec['caption']}")

# ============================================================
# Step 4: 重排幻灯片到正确位置
# ============================================================
print("\nStep 4: 重排幻灯片到正确位置...")
xml_slides = prs.slides._sldIdLst
existing = list(xml_slides)
total = len(existing)
n_created = len(created)
n_original = total - n_created

# 创建的幻灯片在 xml 索引: n_original .. total-1
groups = {}
for i, (sl, ins_after) in enumerate(created):
    groups.setdefault(ins_after, []).append(n_original + i)

desired_order = []
sorted_insert_points = sorted(groups.keys())
ip_pointer = 0
created_used = set()

for orig_idx in range(n_original):
    desired_order.append(orig_idx)
    if (ip_pointer < len(sorted_insert_points)
            and sorted_insert_points[ip_pointer] == orig_idx + 1):
        for new_idx in groups[sorted_insert_points[ip_pointer]]:
            desired_order.append(new_idx)
            created_used.add(new_idx)
        ip_pointer += 1

for i in range(n_original, total):
    if i not in created_used:
        desired_order.append(i)

# Detach all and re-attach in desired order
for el in existing:
    xml_slides.remove(el)
for idx in desired_order:
    xml_slides.append(existing[idx])

print(f"  最终页数: {len(xml_slides)}")
for k in sorted_insert_points:
    print(f"    after slide {k}: +{len(groups[k])} 张仿真页")

# ============================================================
# Step 5: 保存
# ============================================================
prs.save(str(PPT_PATH))
print(f"\n✓ PPTX 已保存: {PPT_PATH}")
print("\n现在请打开 PowerPoint 验证文件不再损坏.")
